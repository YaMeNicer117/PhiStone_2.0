"""Clean rows with only untested activity placeholders.

This script reads a reclassified activity Excel file and removes rows
where the only provided activity test values are explicitly placeholders like
"-", "--", "nt", "nt.", "nd", "nd.".
"""

import argparse
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook, Workbook
except ImportError:
    load_workbook = None
    Workbook = None

# 当前路径与默认文件配置
SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT_FILE = SCRIPT_DIR / "human_mtb_activity_reclassified.xlsx"
DEFAULT_OUTPUT_FILE = SCRIPT_DIR / "human_mtb_activity_reclassified_cleaned.xlsx"

# 需直接删除的无效活性占位符（全部转为小写以进行无大小写区分比对）
TARGET_STRINGS = {"-", "--", "nt", "nt.", "nd", "nd.", "n/t", "n/d", "n.t.", "n.d."}

# 所有的活性数值列表头变体，用于定位数值列
ACTIVITY_VALUE_ALIASES = {
    "mic value", "mic",
    "ic50 value", "ic50",
    "% inhibition", "inhibition", "inhibition value",
    "pic50", "pic50 value",
    "pmic", "pmic value",
    "zoi (mm)", "zoi", "zoi value"
}

def normalize_header(value) -> str:
    """清理并标准化表头以便于精确匹配。"""
    if value is None:
        return ""
    return " ".join(str(value).replace("\ufeff", "").split()).casefold()

def main():
    parser = argparse.ArgumentParser(
        description="清理仅包含未测试/未定活性占位符（如 nt, nd, --）的数据行。"
    )
    parser.add_argument(
        "--input", 
        type=Path, 
        default=DEFAULT_INPUT_FILE,
        help=f"输入 XLSX（默认：{DEFAULT_INPUT_FILE.name}）"
    )
    parser.add_argument(
        "--output", 
        type=Path, 
        default=DEFAULT_OUTPUT_FILE,
        help=f"输出的清理后 XLSX（默认：{DEFAULT_OUTPUT_FILE.name}）"
    )
    args = parser.parse_args()

    if load_workbook is None or Workbook is None:
        sys.exit("[错误] 缺少 openpyxl 库。请在当前环境安装 openpyxl。")

    input_path = args.input.resolve()
    output_path = args.output.resolve()

    if not input_path.exists():
        sys.exit(f"[错误] 找不到输入文件: {input_path}")

    print(f"正在读取文件: {input_path}")
    
    # 仅读取数据，不读取样式
    wb = load_workbook(input_path, read_only=True, data_only=True)
    # 仅使用写入模式生成纯数据、无格式的表格
    wb_out = Workbook(write_only=True)

    total_rows_scanned = 0
    total_rows_deleted = 0
    total_rows_kept = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws_out = wb_out.create_sheet(title=sheet_name)
        
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
            
        # 写入表头
        header_row = rows[0]
        ws_out.append(header_row)
        total_rows_scanned += 1
        
        # 定位所有活性测试数值列的索引
        activity_indices = []
        for idx, cell in enumerate(header_row):
            if normalize_header(cell) in ACTIVITY_VALUE_ALIASES:
                activity_indices.append(idx)
        
        for row in rows[1:]:
            total_rows_scanned += 1
            
            # 如果整行都是空的，直接跳过不作统计
            if not any(val is not None and str(val).strip() for val in row):
                continue

            # 若未找到活性列，为了不丢失数据，默认保留整行
            if not activity_indices:
                ws_out.append(row)
                total_rows_kept += 1
                continue

            has_valid_data = False
            has_target_string = False

            # 遍历该行中属于活性测试的所有单元格
            for idx in activity_indices:
                if idx >= len(row):
                    continue
                val = row[idx]
                if val is None:
                    continue
                    
                s = str(val).strip().lower()
                if not s:  # 跳过纯空字符串
                    continue
                    
                if s in TARGET_STRINGS:
                    has_target_string = True
                else:
                    # 只要存在任何一个不属于占位符的有效数据，就视作有有效活性
                    has_valid_data = True
                    
            # 核心判断逻辑：没有有效数据，且存在 nt/nd/-- 等占位符
            if not has_valid_data and has_target_string:
                total_rows_deleted += 1
                continue
                
            # 不满足删除条件的行，原样追加至新表
            ws_out.append(row)
            total_rows_kept += 1

    # 关闭读取流
    wb.close()

    print("\n处理完成！")
    print(f"总计处理 Sheet 数量: {len(wb.sheetnames)}")
    print(f"扫描总行数 (包含表头): {total_rows_scanned}")
    print(f"移除的无数据/占位符行数: {total_rows_deleted}")
    print(f"最终保留有效数据行数: {total_rows_kept}")
    
    print(f"\n正在保存清理后的文件: {output_path}")
    wb_out.save(output_path)
    print("保存成功！")

if __name__ == "__main__":
    main()
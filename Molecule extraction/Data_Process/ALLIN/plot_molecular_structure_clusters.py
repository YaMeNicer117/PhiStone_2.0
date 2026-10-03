"""读取活性工作簿，对唯一去盐母体进行结构聚类，只输出一张二维分布图。

流程：SMILES 去盐与去重 -> ECFP4 -> 10维 UMAP -> MiniBatchKMeans 粗分组。
二维图从同一份10维表示生成，默认24组，仅用于概览，不表示24种天然化学家族。
仅使用结构，活性值不参与聚类；保留手性，不做电荷或互变异构体标准化。
依赖：openpyxl、rdkit、numpy、umap-learn、scikit-learn、matplotlib。
默认读取脚本同目录的 activity_mic_level.xlsx，原工作簿保持只读。
"""

from __future__ import annotations

import argparse
from itertools import islice
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from rdkit import Chem, DataStructs, RDLogger
from rdkit.Chem import rdFingerprintGenerator
from rdkit.Chem.MolStandardize import rdMolStandardize
from sklearn.cluster import MiniBatchKMeans
from umap import UMAP


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT = SCRIPT_DIR / "activity_mic_level.xlsx"
DEFAULT_OUTPUT = SCRIPT_DIR / "activity_mic_level_structure_clusters.png"

HEADER_SCAN_ROWS = 20
FP_RADIUS = 2
FP_SIZE = 2048
UMAP_N_NEIGHBORS = 100
CLUSTER_DIMENSIONS = 10
N_CLUSTERS = 24
DISPLAY_MIN_DIST = 0.25
RANDOM_SEED = 42

# 离散配色用于区分分组，颜色之间没有数值大小关系。
CLUSTER_COLORS = (
    "#1F77B4", "#E6550D", "#2CA02C", "#D62728", "#9467BD", "#8C564B",
    "#E377C2", "#7F7F7F", "#A6A61B", "#17BECF", "#393B79", "#637939",
    "#8C6D31", "#843C39", "#7B4173", "#3182BD", "#31A354", "#756BB1",
    "#E7BA52", "#D6616B", "#6BAED6", "#74C476", "#CE6DBD", "#9E9AC8",
)


def read_unique_molecules(input_path: Path) -> list[Chem.Mol]:
    """先收集非空 SMILES，再按规范的去盐母体 SMILES 跨 Sheet 去重。"""
    raw_smiles: set[str] = set()
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    matched_sheets = 0
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            smiles_column = None
            for row in islice(rows, HEADER_SCAN_ROWS):
                smiles_column = next(
                    (
                        index
                        for index, value in enumerate(row)
                        if value is not None and str(value).strip().casefold() == "smiles"
                    ),
                    None,
                )
                if smiles_column is not None:
                    break

            if smiles_column is None:
                print(f"跳过 Sheet：{sheet.title}（未找到 SMILES 表头）", flush=True)
                continue

            matched_sheets += 1
            for row in rows:
                if smiles_column >= len(row) or row[smiles_column] is None:
                    continue
                text = str(row[smiles_column]).strip()
                if text and text.casefold() != "smiles":
                    raw_smiles.add(text)
            print(f"已读取 Sheet：{sheet.title}", flush=True)
    finally:
        workbook.close()

    if not matched_sheets:
        raise ValueError(f"所有 Sheet 的前 {HEADER_SCAN_ROWS} 行均未找到 SMILES 列。")

    chooser = rdMolStandardize.LargestFragmentChooser(preferOrganic=True)
    unique_parents: dict[str, Chem.Mol] = {}
    invalid_count = 0
    for smiles in sorted(raw_smiles):
        try:
            molecule = Chem.MolFromSmiles(smiles)
            if molecule is None or molecule.GetNumAtoms() == 0:
                invalid_count += 1
                continue
            parent = chooser.choose(molecule)
            parent_smiles = Chem.MolToSmiles(parent, canonical=True, isomericSmiles=True)
        except (ValueError, RuntimeError):
            invalid_count += 1
            continue
        unique_parents.setdefault(parent_smiles, parent)

    print(
        f"不同原始 SMILES：{len(raw_smiles):,}；无效：{invalid_count:,}；"
        f"唯一去盐母体：{len(unique_parents):,}",
        flush=True,
    )
    if len(unique_parents) < 3:
        raise ValueError("至少需要 3 个有效且不重复的母体结构。")
    # 稳定顺序配合固定随机种子，减少 Excel 行顺序变化对结果的影响。
    return [unique_parents[smiles] for smiles in sorted(unique_parents)]


def make_fingerprint_matrix(molecules: list[Chem.Mol]) -> np.ndarray:
    generator = rdFingerprintGenerator.GetMorganGenerator(
        radius=FP_RADIUS, fpSize=FP_SIZE, includeChirality=True
    )
    matrix = np.zeros((len(molecules), FP_SIZE), dtype=np.uint8)
    for index, molecule in enumerate(molecules):
        fingerprint = generator.GetFingerprint(molecule)
        DataStructs.ConvertToNumpyArray(fingerprint, matrix[index])
    return matrix


def cluster_and_embed(matrix: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    """在10维表示中粗分组，再将同一表示投影到二维；分组标签不参与投影。"""
    common_parameters = dict(
        n_neighbors=min(UMAP_N_NEIGHBORS, len(matrix) - 1),
        random_state=RANDOM_SEED,
        init="random",
        n_jobs=1,
        low_memory=True,
    )
    print(f"计算 {CLUSTER_DIMENSIONS} 维 UMAP 聚类表示……", flush=True)
    cluster_embedding = UMAP(
        n_components=CLUSTER_DIMENSIONS, min_dist=0.0,
        metric="jaccard", **common_parameters,
    ).fit_transform(matrix)

    n_clusters = min(N_CLUSTERS, len(matrix))
    print(f"执行 MiniBatchKMeans 结构粗分组，目标组数：{n_clusters}……", flush=True)
    labels = MiniBatchKMeans(
        n_clusters=n_clusters,
        batch_size=2048,
        n_init=10,
        random_state=RANDOM_SEED,
    ).fit_predict(cluster_embedding)

    print("从10维结构表示计算二维 UMAP 展示坐标……", flush=True)
    coordinates = UMAP(
        n_components=2, min_dist=DISPLAY_MIN_DIST,
        metric="euclidean", **common_parameters,
    ).fit_transform(cluster_embedding)
    return coordinates, labels


def save_plot(coordinates: np.ndarray, labels: np.ndarray, output_path: Path) -> None:
    cluster_ids, counts = np.unique(labels, return_counts=True)
    order = np.argsort(-counts, kind="stable")
    cluster_ids, counts = cluster_ids[order], counts[order]
    cluster_count = len(cluster_ids)
    point_size = float(np.clip(15000 / len(labels), 2, 16))
    palette = (
        CLUSTER_COLORS[:cluster_count]
        if cluster_count <= len(CLUSTER_COLORS)
        else plt.get_cmap("turbo")(np.linspace(0.03, 0.97, cluster_count))
    )

    figure, axis = plt.subplots(figsize=(13, 9))
    try:
        for index, cluster_id in enumerate(cluster_ids):
            mask = labels == cluster_id
            axis.scatter(
                coordinates[mask, 0], coordinates[mask, 1],
                s=point_size, color=palette[index], alpha=0.55, linewidths=0,
                label=f"C{index + 1}  (n={counts[index]:,})",
            )
            if cluster_count <= 30:
                center = np.median(coordinates[mask], axis=0)
                axis.annotate(
                    f"C{index + 1}", center, ha="center", va="center", fontsize=8,
                    bbox=dict(boxstyle="round,pad=0.15", facecolor="white", alpha=0.75, edgecolor="none"),
                )

        axis.set_title(
            "Molecular structural overview\n"
            f"{len(labels):,} unique parents | {cluster_count} coarse groups",
            fontsize=14, pad=14,
        )
        axis.set_xlabel("UMAP 1")
        axis.set_ylabel("UMAP 2")
        axis.set_aspect("equal", adjustable="datalim")
        axis.spines[["top", "right"]].set_visible(False)
        axis.legend(
            title="Groups by size", loc="center left", bbox_to_anchor=(1.02, 0.5),
            frameon=False, markerscale=3, fontsize=8, labelspacing=0.65,
        )
        figure.text(
            0.5, 0.02,
            f"ECFP4 / Jaccard | MiniBatchKMeans on {CLUSTER_DIMENSIONS}D UMAP | "
            "Coarse groups; 2D distances are approximate",
            ha="center", fontsize=9, color="#555555",
        )
        figure.tight_layout(rect=(0, 0.04, 1, 1))
        output_path.parent.mkdir(parents=True, exist_ok=True)
        figure.savefig(output_path, dpi=300, bbox_inches="tight", facecolor="white")
    finally:
        plt.close(figure)
    print(f"结构粗分组：{cluster_count}；已分组结构：{len(labels):,}", flush=True)
    print(f"图片已保存：{output_path.resolve()}", flush=True)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="输入 Excel 路径")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="输出 PNG 路径")
    arguments = parser.parse_args()
    if arguments.output.suffix.lower() != ".png":
        parser.error("输出文件须使用 .png 扩展名。")

    RDLogger.DisableLog("rdApp.error")
    print(f"读取工作簿：{arguments.input}", flush=True)
    molecules = read_unique_molecules(arguments.input)
    print("生成 ECFP4 指纹……", flush=True)
    matrix = make_fingerprint_matrix(molecules)
    coordinates, labels = cluster_and_embed(matrix)
    save_plot(coordinates, labels, arguments.output)


if __name__ == "__main__":
    main()

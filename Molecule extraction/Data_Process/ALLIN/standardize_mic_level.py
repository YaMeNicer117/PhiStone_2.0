"""按 IC50、抑制率和 MIC 的既定优先级分配项目自定义活性等级。

1. 明确 IC50；2. (0, 25] ug/mL 下的抑制率；3. 明确 MIC；
P4 不再使用；
5. IC50 不等式优先于 MIC 不等式；6. 既有无活性关键词赋 0。

IC50 和 MIC 共用解析规则：范围取中点，丢弃 ± 或 +/- 后的误差，
≈ 视为等于，~ / about / approx 使用数值。
IC50/MIC 标准化为 uM，测试浓度标准化为 ug/mL；抑制率不等号忽略。
不推算 at MIC 或 MIC 倍数浓度，也不把 IC50 换算成 MIC。

仅输出来源、化合物编号、SMILES、标准化 IC50/MIC、原始抑制率及测试浓度、
活性等级和评分依据。使用默认表格外观，保留原始字段的数字格式。
已评分行保持原始顺序且不标色，
无法评分但仍有实际信息的行标黄并放到 Sheet 末尾；仅为空、占位码或
not determined 的记录舍弃。关键词评分不伪造 MIC/IC50 浓度。输入工作簿永不修改。
"""

from __future__ import annotations

import argparse
import math
import os
import re
import sys
import tempfile
import unicodedata
from contextlib import contextmanager
from dataclasses import dataclass, field
from numbers import Real
from pathlib import Path
from threading import Event, Thread
from time import monotonic
from typing import Any, Literal, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from rdkit import Chem, RDLogger
from rdkit.Chem import Descriptors


DEFAULT_INPUT = Path("human_mtb_activity_reclassified_cleaned.xlsx")
DEFAULT_OUTPUT = Path("activity_mic_level.xlsx")

HEADER_SCAN_ROWS = 20
ACTIVITY_LEVEL_BANDS = (
    (2.0, 1.0),
    (12.5, 0.8),
    (32.0, 0.6),
    (64.0, 0.4),
)
INACTIVE_MIC_UM = 256.0
LOW_ACTIVITY_LEVEL = 0.1
INACTIVE_ACTIVITY_LEVEL = 0.0
IC50_LEVEL_BANDS = ((1.0, 1.0), (10.0, 0.75), (25.0, 0.5), (60.0, 0.25), (120.0, 0.125))
# 各区域上界包含本身；评分阈值按从高到低排列，下界包含本身。
INHIBITION_LEVEL_BANDS = (
    (3.125, "A", ((90.0, 1.0), (75.0, 0.75), (50.0, 0.5), (25.0, 0.25))),
    (6.25, "B", ((90.0, 1.0), (75.0, 0.75), (55.0, 0.5), (40.0, 0.25))),
    (12.5, "C", ((95.0, 1.0), (80.0, 0.75), (60.0, 0.5), (45.0, 0.25))),
    (25.0, "D", ((90.0, 0.75), (80.0, 0.5), (65.0, 0.25), (50.0, 0.125))),
)

PRESERVED_COLUMNS = (
    "Source Filename",
    "Compound ID",
    "SMILES",
    "% Inhibition",
    "Inhibition Conc",
)
REQUIRED_INPUT_COLUMNS = PRESERVED_COLUMNS + (
    "MIC Value",
    "MIC Unit",
    "IC50 Value",
    "IC50 Unit",
)
OUTPUT_COLUMNS = (
    "Source Filename",
    "Compound ID",
    "SMILES",
    "IC50 (μM)",
    "MIC (μM)",
    "% Inhibition",
    "Inhibition Conc",
    "Activity Level",
    "Activity Basis",
)

YELLOW_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")

Comparator = Literal["eq", "lt", "le", "gt", "ge"]
UnitDimension = Literal["molar", "mass"]

COMPARATOR_SYMBOLS: dict[Comparator, str] = {
    "eq": "",
    "lt": "<",
    "le": "<=",
    "gt": ">",
    "ge": ">=",
}


class ProcessingError(RuntimeError):
    """Raised when workbook validation or processing cannot safely continue."""


@dataclass(frozen=True)
class UnitSpec:
    dimension: UnitDimension
    multiplier: float
    canonical: str


@dataclass(frozen=True)
class ParsedMic:
    numeric_value: float | None
    comparator: Comparator = "eq"
    approximate: bool = False
    range_used: bool = False
    keyword_inactive: bool = False

    @property
    def is_numeric(self) -> bool:
        return self.numeric_value is not None


@dataclass(frozen=True)
class SheetSpec:
    name: str
    header_row: int
    columns: dict[str, int]


@dataclass(frozen=True)
class SmilesAnalysis:
    valid: bool
    molecular_weight: float | None
    total_charge: int | None


@dataclass(frozen=True)
class StandardizedMeasurement:
    parsed: ParsedMic
    unit: UnitSpec | None
    micromolar: float | None
    reason: str = ""


@dataclass(frozen=True)
class InhibitionMeasurement:
    parsed: ParsedMic
    percent: float | None
    concentration: float | None
    unit: UnitSpec | None
    reason: str = ""


@dataclass(frozen=True)
class ActivityDecision:
    level: float | None
    priority: int = 0
    source: str = ""
    detail: str = ""


@dataclass(frozen=True)
class OutputRow:
    values: list[Any]
    source_row: int
    yellow: bool


@dataclass
class SheetStats:
    name: str
    source_rows: int = 0
    removed_blank: int = 0
    output_rows: int = 0
    calculated: int = 0
    keyword_rows: int = 0
    excluded_keyword_rows: int = 0
    range_rows: int = 0
    manual_rows: int = 0
    priority_counts: dict[int, int] = field(default_factory=dict)


def progress_rows(rows, total: int, label: str):
    """按已完成行数显示进度，每 0.25 秒最多刷新一次。"""
    started = monotonic()
    last_refresh = started
    completed = 0
    previous_width = 0

    def display():
        nonlocal previous_width
        elapsed = monotonic() - started
        fraction = completed / total if total else 1.0
        filled = int(24 * fraction)
        remaining = f"{elapsed / completed * (total - completed):.0f}s" if completed else "--"
        message = (
            f"{label} [{'#' * filled}{'-' * (24 - filled)}] "
            f"{fraction:6.1%} {completed}/{total} 行 "
            f"已用 {elapsed:.0f}s 剩余约 {remaining}"
        )
        print("\r" + message.ljust(previous_width), end="", file=sys.stderr, flush=True)
        previous_width = len(message)

    display()
    try:
        for row in rows:
            yield row
            completed += 1
            now = monotonic()
            if now - last_refresh >= 0.25:
                display()
                last_refresh = now
    finally:
        display()
        print(file=sys.stderr, flush=True)


@contextmanager
def timed_stage(label: str):
    """读取、预检、保存无法取得准确百分比，仅显示动态状态和耗时。"""
    started = monotonic()
    stopped = Event()

    def animate():
        index = 0
        while not stopped.wait(0.25):
            symbol = "|/-\\"[index % 4]
            print(
                f"\r[{symbol}] {label}，已用 {monotonic() - started:.0f}s",
                end="", file=sys.stderr, flush=True,
            )
            index += 1

    print(f"[开始] {label}", file=sys.stderr, flush=True)
    worker = Thread(target=animate, daemon=True)
    worker.start()
    status = "中断"
    try:
        yield
        status = "完成"
    finally:
        stopped.set()
        worker.join()
        print(
            f"\r[{status}] {label}，耗时 {monotonic() - started:.1f}s",
            file=sys.stderr, flush=True,
        )


def _unit_aliases() -> dict[str, UnitSpec]:
    aliases: dict[str, UnitSpec] = {}

    def add(
        names: Sequence[str],
        dimension: UnitDimension,
        multiplier: float,
        canonical: str,
    ) -> None:
        spec = UnitSpec(dimension, multiplier, canonical)
        for name in names:
            aliases[name] = spec

    # For molar units, multiplier converts the source value directly to uM.
    add(("m", "mol/l", "moll-1", "molar"), "molar", 1_000_000.0, "M")
    add(
        ("mm", "mmol", "mmol/l", "mmoll-1", "millimolar"),
        "molar",
        1_000.0,
        "mM",
    )
    add(
        (
            "um",
            "um/ml",
            "umol/l",
            "umoll-1",
            "micromolar",
            "micromol/l",
            "micromoll-1",
        ),
        "molar",
        1.0,
        "uM",
    )
    # uM/ml is handled above as a documented typo for uM. In contrast,
    # micromoles per mL is a true amount-per-volume unit equal to 1000 uM.
    add(("umol/ml", "umolml-1"), "molar", 1_000.0, "umol/mL")
    add(
        ("nm", "nmol/l", "nmoll-1", "nanomolar"),
        "molar",
        0.001,
        "nM",
    )
    add(
        ("pm", "pmol/l", "pmoll-1", "picomolar"),
        "molar",
        0.000001,
        "pM",
    )

    # For mass units, multiplier converts one source unit to g/L.
    add(("g/l", "gl-1", "gram/l", "grams/l"), "mass", 1.0, "g/L")
    add(("mg/ml", "mgml-1"), "mass", 1.0, "mg/mL")
    add(("mg/l", "mgl-1"), "mass", 0.001, "mg/L")
    add(
        ("ug/ml", "ugml-1", "mcg/ml", "mcgml-1"),
        "mass",
        0.001,
        "ug/mL",
    )
    add(
        ("ug/l", "ugl-1", "mcg/l", "mcgl-1"),
        "mass",
        0.000001,
        "ug/L",
    )
    add(("ng/ml", "ngml-1"), "mass", 0.000001, "ng/mL")
    return aliases


UNIT_ALIASES = _unit_aliases()

UNSIGNED_NUMBER_PATTERN = (
    r"(?:(?:\d+(?:[.,]\d*)?)|(?:[.,]\d+))(?:[eE][+-]?\d+)?"
)
NUMBER_RE = re.compile(rf"^[+-]?{UNSIGNED_NUMBER_PATTERN}")
RANGE_RE = re.compile(
    rf"^(?P<start>{UNSIGNED_NUMBER_PATTERN})"
    rf"(?:->|→|~|to|-)"
    rf"(?P<end>{UNSIGNED_NUMBER_PATTERN})",
    re.IGNORECASE,
)
MIC_PREFIX_RE = re.compile(r"^(?:mic|ic50)(?:value)?(?:[:=])?", re.IGNORECASE)
APPROX_PREFIX_RE = re.compile(
    r"^(?:~|about|approx(?:imately)?\.?)",
    re.IGNORECASE,
)
COMPARATOR_PREFIX_RE = re.compile(
    r"^(?:"
    r"(?P<le><=|=<|≤|lessthanorequalto|nomorethan|atmost)|"
    r"(?P<ge>>=|=>|≥|greaterthanorequalto|morethanorequalto|atleast)|"
    r"(?P<lt>(?:<<|<|lessthan|below|under))|"
    r"(?P<gt>(?:>>|>|greaterthan|morethan|above|over))|"
    r"(?P<eq>=|≈)"
    r")",
    re.IGNORECASE,
)

SHORT_INACTIVE_PATTERNS = (
    re.compile(r"(?<![a-z0-9])na(?![a-z0-9])", re.IGNORECASE),
    re.compile(r"(?<![a-z0-9])n\s*[/\.]\s*a\.?(?![a-z0-9])", re.IGNORECASE),
)
EXCLUDED_MIC_RE = re.compile(
    r"(?:\*+|-{2,}|n[/\.]?t\.?|n[/\.]?d\.?|nma|res|notdetermined)", re.IGNORECASE,
)
INACTIVE_KEYWORDS_COMPACT = (
    "inactive",
    "notactive",
    "noactivity",
    "noinhibition",
    "poor",
    "resistant",
    "weak",
    "lowactivity",
    "noeffective",
    "less",
    "lack",
    "loss",
    "ineffective",
)


def normalize_text(value: Any) -> str:
    """Return normalized text without changing its semantic content."""

    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def normalize_header(value: Any) -> str:
    text = normalize_text(value)
    return re.sub(r"\s+", " ", text).casefold()


def normalize_unit_token(value: Any) -> str:
    """Normalize common Unicode and textual unit spellings for lookup."""

    text = normalize_text(value).casefold()
    if not text:
        return ""

    text = text.replace("μ", "u").replace("µ", "u")
    text = text.replace("−", "-").replace("–", "-").replace("—", "-")
    text = text.replace("⁻", "-").replace("¹", "1")
    text = text.replace("·", "").replace("⋅", "").replace("*", "")
    text = text.replace("\\", "/").replace("per", "/")
    text = text.replace("^", "").replace(".", "")
    text = re.sub(r"\s+", "", text)
    return text.strip("()[]{};,:")


def parse_unit(value: Any) -> UnitSpec | None:
    return UNIT_ALIASES.get(normalize_unit_token(value))


def is_blank_mic(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _discard_uncertainty(text: str) -> str:
    positions = [position for marker in ("±", "+/-") if (position := text.find(marker)) >= 0]
    if not positions:
        return text
    return text[: min(positions)]


def _consume_approximation(text: str) -> tuple[str, bool]:
    match = APPROX_PREFIX_RE.match(text)
    if not match:
        return text, False
    return text[match.end() :], True


def _consume_comparator(text: str) -> tuple[str, Comparator]:
    match = COMPARATOR_PREFIX_RE.match(text)
    if not match:
        return text, "eq"

    comparator: Comparator = "eq"
    for candidate in ("le", "ge", "lt", "gt", "eq"):
        if match.group(candidate) is not None:
            comparator = candidate  # type: ignore[assignment]
            break
    return text[match.end() :], comparator


def _parse_numeric_token(token: str) -> float | None:
    mantissa = token
    exponent = ""
    exponent_match = re.search(r"[eE][+-]?\d+$", token)
    if exponent_match:
        mantissa = token[: exponent_match.start()]
        exponent = exponent_match.group(0)

    if "," in mantissa and "." in mantissa:
        return None
    mantissa = mantissa.replace(",", ".")
    try:
        result = float(mantissa + exponent)
    except ValueError:
        return None
    return result if math.isfinite(result) else None


def _numeric_string_parse(value: str, allow_unit_suffix: bool = True) -> ParsedMic | None:
    text = re.sub(r"\s+", "", normalize_text(value))
    text = text.replace("−", "-").replace("–", "-").replace("—", "-")
    text = _discard_uncertainty(text)
    text = MIC_PREFIX_RE.sub("", text, count=1).strip()

    text, approx_before = _consume_approximation(text)
    text, comparator = _consume_comparator(text)
    text, approx_after = _consume_approximation(text)
    approximate = approx_before or approx_after

    range_match = RANGE_RE.match(text)
    if range_match:
        start_value = _parse_numeric_token(range_match.group("start"))
        end_value = _parse_numeric_token(range_match.group("end"))
        if start_value is None or end_value is None:
            return None
        if start_value <= 0 or end_value <= 0:
            return None

        suffix = text[range_match.end() :].strip().rstrip(";,.").strip()
        if suffix:
            stripped_suffix = suffix.strip("()[]{}")
            if not allow_unit_suffix or parse_unit(stripped_suffix) is None:
                return None

        return ParsedMic(
            numeric_value=(start_value + end_value) / 2.0,
            comparator=comparator,
            approximate=approximate,
            range_used=True,
        )

    match = NUMBER_RE.match(text)
    if not match:
        return None

    numeric_value = _parse_numeric_token(match.group(0))
    if numeric_value is None:
        return None

    suffix = text[match.end() :].strip().rstrip(";,.").strip()
    if suffix:
        stripped_suffix = suffix.strip("()[]{}")
        if not allow_unit_suffix or parse_unit(stripped_suffix) is None:
            return None

    return ParsedMic(
        numeric_value=numeric_value,
        comparator=comparator,
        approximate=approximate,
    )


def contains_inactive_keyword(value: Any) -> bool:
    text = normalize_text(value)
    if not text:
        return False
    if any(pattern.search(text) for pattern in SHORT_INACTIVE_PATTERNS):
        return True
    compact = re.sub(r"\s+", "", text).casefold()
    return any(keyword in compact for keyword in INACTIVE_KEYWORDS_COMPACT)


def contains_excluded_mic_keyword(value: Any) -> bool:
    """仅识别整个字段为占位码的情况，不因数字旁有 * 而删除实际信息。"""

    text = re.sub(r"\s+", "", normalize_text(value))
    if not text:
        return False
    return EXCLUDED_MIC_RE.fullmatch(text) is not None


def parse_mic_value(value: Any, allow_unit_suffix: bool = True) -> ParsedMic:
    """IC50 与 MIC 共用解析器，合法数值优先于无活性关键词。"""

    if isinstance(value, Real) and not isinstance(value, bool):
        numeric_value = float(value)
        if math.isfinite(numeric_value):
            return ParsedMic(numeric_value=numeric_value)
        return ParsedMic(numeric_value=None)

    if isinstance(value, str):
        parsed_numeric = _numeric_string_parse(value, allow_unit_suffix)
        if parsed_numeric is not None:
            return parsed_numeric
        if contains_inactive_keyword(value):
            return ParsedMic(numeric_value=None, comparator="ge", keyword_inactive=True)

    return ParsedMic(numeric_value=None)


def analyze_smiles(
    smiles_value: Any,
    cache: dict[str, SmilesAnalysis],
) -> SmilesAnalysis:
    """Return validity, average MolWt, and total formal charge for a SMILES."""

    smiles = normalize_text(smiles_value)
    if not smiles:
        return SmilesAnalysis(False, None, None)
    if smiles in cache:
        return cache[smiles]

    try:
        molecule = Chem.MolFromSmiles(smiles)
        if molecule is None:
            result = SmilesAnalysis(False, None, None)
        else:
            molecular_weight = float(Descriptors.MolWt(molecule))
            if not math.isfinite(molecular_weight) or molecular_weight <= 0:
                result = SmilesAnalysis(False, None, None)
            else:
                total_charge = int(Chem.GetFormalCharge(molecule))
                result = SmilesAnalysis(True, molecular_weight, total_charge)
    except Exception:
        result = SmilesAnalysis(False, None, None)

    cache[smiles] = result
    return result


def convert_to_micromolar(
    numeric_value: float,
    unit: UnitSpec,
    molecular_weight: float | None,
) -> float | None:
    if not math.isfinite(numeric_value) or numeric_value <= 0:
        return None

    if unit.dimension == "molar":
        micromolar = numeric_value * unit.multiplier
    else:
        if molecular_weight is None or molecular_weight <= 0:
            return None
        grams_per_liter = numeric_value * unit.multiplier
        micromolar = grams_per_liter * 1_000_000.0 / molecular_weight

    if not math.isfinite(micromolar) or micromolar <= 0:
        return None
    return micromolar


def calculate_activity_level(
    micromolar: float,
    comparator: Comparator = "eq",
) -> float:
    """Assign one fixed activity level from the converted MIC boundary."""

    effective_mic = micromolar
    if comparator == "gt":
        effective_mic = math.nextafter(micromolar, math.inf)
    elif comparator == "lt":
        effective_mic = math.nextafter(micromolar, -math.inf)

    for upper_bound, activity_level in ACTIVITY_LEVEL_BANDS:
        if effective_mic <= upper_bound:
            return activity_level
    if effective_mic < 128.0:
        return 0.2
    if effective_mic < INACTIVE_MIC_UM:
        return LOW_ACTIVITY_LEVEL
    return INACTIVE_ACTIVITY_LEVEL


def calculate_ic50_activity_level(
    micromolar: float,
    comparator: Comparator = "eq",
) -> float:
    if comparator == "gt":
        micromolar = math.nextafter(micromolar, math.inf)
    elif comparator == "lt":
        micromolar = math.nextafter(micromolar, -math.inf)
    for upper_bound, level in IC50_LEVEL_BANDS:
        if micromolar <= upper_bound:
            return level
    return 0.0


def standardize_measurement(
    value: Any, unit_value: Any, smiles: SmilesAnalysis,
) -> StandardizedMeasurement:
    parsed = parse_mic_value(value)
    unit = parse_unit(unit_value)
    if not parsed.is_numeric:
        return StandardizedMeasurement(parsed, unit, None, "缺少可解析数值")
    if unit is None:
        return StandardizedMeasurement(parsed, unit, None, "单位缺失或无法识别")
    if unit.dimension == "mass" and not smiles.valid:
        return StandardizedMeasurement(parsed, unit, None, "质量浓度换算缺少有效分子量")
    micromolar = convert_to_micromolar(parsed.numeric_value, unit, smiles.molecular_weight)
    reason = "" if micromolar is not None else "数值非正或换算失败"
    return StandardizedMeasurement(parsed, unit, micromolar, reason)


def parse_inhibition_concentration(
    value: Any, smiles: SmilesAnalysis,
) -> tuple[float | None, UnitSpec | None, str]:
    """只处理带单位的明确浓度；MIC 倍数、范围和不等式不参与换算。"""
    text = re.sub(r"\s+", "", normalize_text(value))
    text = re.sub(r"^at", "", text, count=1, flags=re.IGNORECASE)
    text = re.sub(r"^[=≈]", "", text, count=1)
    match = NUMBER_RE.match(text)
    if match is None:
        return None, None, "测试浓度缺失或不是明确数值和单位"
    number = _parse_numeric_token(match.group(0))
    unit = parse_unit(text[match.end():])
    if number is None or number <= 0 or unit is None:
        return None, unit, "测试浓度非正、单位缺失或表达不明确"
    if unit.dimension == "mass":
        concentration = number * unit.multiplier * 1000.0
    elif smiles.valid:
        concentration = number * unit.multiplier * smiles.molecular_weight / 1000.0
    else:
        return None, unit, "测试浓度从摩尔单位换算需要有效分子量"
    if not math.isfinite(concentration) or concentration <= 0:
        return None, unit, "测试浓度换算失败"
    return concentration, unit, ""


def standardize_inhibition(
    value: Any, number_format: str, concentration_value: Any, smiles: SmilesAnalysis,
) -> InhibitionMeasurement:
    # Excel 的百分比数值单元格可能存为 0.9，显示为 90%；普通数值仍按原值。
    if isinstance(value, Real) and not isinstance(value, bool) and "%" in number_format:
        value = float(value) * 100.0
    elif isinstance(value, str):
        value = normalize_text(value).rstrip("%").strip()
    parsed = parse_mic_value(value, allow_unit_suffix=False)
    percent = parsed.numeric_value
    reasons = []
    if percent is None or not 0.0 <= percent <= 100.0:
        percent = None
        reasons.append("抑制率缺失、无法解析或不在0至100之间")
    concentration, unit, reason = parse_inhibition_concentration(concentration_value, smiles)
    if reason:
        reasons.append(reason)
    elif concentration > 25.0:
        reasons.append("测试浓度超过25 μg/mL，不参与自动评分")
    return InhibitionMeasurement(parsed, percent, concentration, unit, "；".join(reasons))


def calculate_inhibition_activity_level(percent: float, concentration: float) -> tuple[float, str]:
    for upper_bound, region, thresholds in INHIBITION_LEVEL_BANDS:
        if concentration <= upper_bound:
            for lower_bound, level in thresholds:
                if percent >= lower_bound:
                    return level, region
            return 0.0, region
    raise ValueError("抑制率评分浓度必须在 (0, 25] μg/mL 内。")


def select_activity(
    ic50: StandardizedMeasurement,
    mic: StandardizedMeasurement,
    inhibition: InhibitionMeasurement,
) -> ActivityDecision:
    """严格按已启用优先级选择，0 分也是成功结果，不再尝试后续指标。"""
    if ic50.micromolar is not None and ic50.parsed.comparator == "eq":
        return ActivityDecision(calculate_ic50_activity_level(ic50.micromolar), 1, "IC50")

    concentration = inhibition.concentration
    inhibition_usable = (
        inhibition.percent is not None and concentration is not None
        and 0.0 < concentration <= 25.0
    )
    if inhibition_usable:
        level, region = calculate_inhibition_activity_level(inhibition.percent, concentration)
        return ActivityDecision(level, 2, "% Inhibition", f"区域{region}")

    if mic.micromolar is not None and mic.parsed.comparator == "eq":
        return ActivityDecision(calculate_activity_level(mic.micromolar), 3, "MIC")

    for name, measurement, calculate in (
        ("IC50", ic50, calculate_ic50_activity_level),
        ("MIC", mic, calculate_activity_level),
    ):
        if measurement.micromolar is not None and measurement.parsed.comparator != "eq":
            return ActivityDecision(
                calculate(measurement.micromolar, measurement.parsed.comparator),
                5, name, "不等式边界评分",
            )

    for name, parsed in (("IC50", ic50.parsed), ("MIC", mic.parsed), ("% Inhibition", inhibition.parsed)):
        if parsed.keyword_inactive:
            return ActivityDecision(0.0, 6, name, "无活性关键词")
    return ActivityDecision(None)


def format_standardized_mic_value(
    micromolar: float,
    comparator: Comparator,
) -> float | str:
    """Keep a normalized inequality comparator with the converted uM value."""

    if comparator == "eq":
        return micromolar
    return f"{COMPARATOR_SYMBOLS[comparator]}{micromolar:.15g}"


def _header_positions(sheet: Worksheet, row_number: int) -> dict[str, list[int]]:
    positions: dict[str, list[int]] = {}
    for column_number in range(1, sheet.max_column + 1):
        key = normalize_header(sheet.cell(row=row_number, column=column_number).value)
        if key:
            positions.setdefault(key, []).append(column_number)
    return positions


def find_header(sheet: Worksheet) -> tuple[int, dict[str, list[int]]] | None:
    mic_key = normalize_header("MIC Value")
    last_scan_row = min(HEADER_SCAN_ROWS, sheet.max_row)
    for row_number in range(1, last_scan_row + 1):
        positions = _header_positions(sheet, row_number)
        if mic_key in positions:
            return row_number, positions
    return None


def preflight_workbook(workbook: Any) -> tuple[list[SheetSpec], list[str]]:
    specs: list[SheetSpec] = []
    skipped: list[str] = []
    errors: list[str] = []

    for sheet in workbook.worksheets:
        header = find_header(sheet)
        if header is None:
            skipped.append(sheet.title)
            continue

        header_row, positions = header
        columns: dict[str, int] = {}
        sheet_errors: list[str] = []
        for column_name in REQUIRED_INPUT_COLUMNS:
            matches = positions.get(normalize_header(column_name), [])
            if not matches:
                sheet_errors.append(f"缺少列 {column_name!r}")
            elif len(matches) > 1:
                sheet_errors.append(f"列 {column_name!r} 重复 {len(matches)} 次")
            else:
                columns[column_name] = matches[0]

        if sheet_errors:
            errors.append(f"Sheet {sheet.title!r}: " + "；".join(sheet_errors))
        else:
            specs.append(SheetSpec(sheet.title, header_row, columns))

    if errors:
        raise ProcessingError("工作簿预检失败：\n- " + "\n- ".join(errors))
    if not specs:
        raise ProcessingError(
            f"前 {HEADER_SCAN_ROWS} 行内未找到包含 'MIC Value' 的数据 Sheet。"
        )
    return specs, skipped


def apply_row_fill(sheet: Worksheet, row_number: int, fill: PatternFill) -> None:
    for column_number in range(1, len(OUTPUT_COLUMNS) + 1):
        sheet.cell(row=row_number, column=column_number).fill = fill


def copy_output_values(
    source_sheet: Worksheet,
    source_row: int,
    spec: SheetSpec,
    ic50: StandardizedMeasurement,
    mic: StandardizedMeasurement,
    decision: ActivityDecision,
) -> list[Any]:
    # 按 OUTPUT_COLUMNS 顺序输出；抑制率及测试浓度保持原始内容。
    values = [
        source_sheet.cell(source_row, spec.columns[name]).value
        for name in ("Source Filename", "Compound ID", "SMILES")
    ]
    for measurement in (ic50, mic):
        values.append(
            format_standardized_mic_value(measurement.micromolar, measurement.parsed.comparator)
            if measurement.micromolar is not None else None
        )
    basis = (
        f"P{decision.priority} {decision.source}"
        if decision.level is not None else "待人工判断"
    )
    if decision.detail:
        basis += f"；{decision.detail}"
    values.extend(source_sheet.cell(source_row, spec.columns[name]).value
                  for name in ("% Inhibition", "Inhibition Conc"))
    values.extend((decision.level, basis))
    return values


def process_sheet(
    source_sheet: Worksheet,
    output_sheet: Worksheet,
    spec: SheetSpec,
    smiles_cache: dict[str, SmilesAnalysis],
) -> SheetStats:
    stats = SheetStats(
        name=source_sheet.title,
        source_rows=max(source_sheet.max_row - spec.header_row, 0),
    )
    output_sheet.append(list(OUTPUT_COLUMNS))

    regular_rows: list[OutputRow] = []
    manual_rows: list[OutputRow] = []

    for source_row in progress_rows(
        range(spec.header_row + 1, source_sheet.max_row + 1),
        stats.source_rows, f"[{spec.name}] 解析与评分",
    ):
        mic_value = source_sheet.cell(source_row, spec.columns["MIC Value"]).value
        ic50_value = source_sheet.cell(source_row, spec.columns["IC50 Value"]).value
        inhibition_cell = source_sheet.cell(source_row, spec.columns["% Inhibition"])
        activity_values = (ic50_value, mic_value, inhibition_cell.value)
        if all(is_blank_mic(value) for value in activity_values):
            stats.removed_blank += 1
            continue
        if all(is_blank_mic(value) or contains_excluded_mic_keyword(value) for value in activity_values):
            stats.excluded_keyword_rows += 1
            continue

        smiles_analysis = analyze_smiles(
            source_sheet.cell(source_row, spec.columns["SMILES"]).value, smiles_cache,
        )
        ic50 = standardize_measurement(
            ic50_value, source_sheet.cell(source_row, spec.columns["IC50 Unit"]).value,
            smiles_analysis,
        )
        mic = standardize_measurement(
            mic_value, source_sheet.cell(source_row, spec.columns["MIC Unit"]).value,
            smiles_analysis,
        )
        inhibition = standardize_inhibition(
            inhibition_cell.value, inhibition_cell.number_format,
            source_sheet.cell(source_row, spec.columns["Inhibition Conc"]).value,
            smiles_analysis,
        )
        decision = select_activity(ic50, mic, inhibition)
        selected = {"IC50": ic50, "MIC": mic, "% Inhibition": inhibition}.get(decision.source)
        needs_manual = decision.level is None

        if needs_manual:
            stats.manual_rows += 1
        else:
            stats.calculated += 1
            stats.priority_counts[decision.priority] = stats.priority_counts.get(decision.priority, 0) + 1
            if decision.priority == 6:
                stats.keyword_rows += 1
            elif selected is not None:
                if selected.parsed.range_used:
                    stats.range_rows += 1

        output_row_data = OutputRow(
            values=copy_output_values(
                source_sheet, source_row, spec, ic50, mic, decision,
            ),
            source_row=source_row,
            yellow=needs_manual,
        )
        if needs_manual:
            manual_rows.append(output_row_data)
        else:
            regular_rows.append(output_row_data)

    for output_row, output_row_data in progress_rows(
        enumerate(regular_rows + manual_rows, start=2),
        len(regular_rows) + len(manual_rows), f"[{spec.name}] 写入结果",
    ):
        output_sheet.append(output_row_data.values)
        stats.output_rows += 1

        # 保留原始百分比等数字格式，避免原来显示 90% 的单元格变成 0.9。
        for name in PRESERVED_COLUMNS:
            column = OUTPUT_COLUMNS.index(name) + 1
            output_sheet.cell(output_row, column).number_format = source_sheet.cell(
                output_row_data.source_row, spec.columns[name],
            ).number_format

        if output_row_data.yellow:
            apply_row_fill(output_sheet, output_row, YELLOW_FILL)

    return stats


def _same_path(left: Path, right: Path) -> bool:
    left_normalized = os.path.normcase(str(left.resolve(strict=False)))
    right_normalized = os.path.normcase(str(right.resolve(strict=False)))
    return left_normalized == right_normalized


def save_workbook_atomically(
    workbook: Workbook,
    output_path: Path,
    overwrite: bool,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output_path.stem}.",
            suffix=".tmp.xlsx",
            dir=output_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        if output_path.exists() and not overwrite:
            raise ProcessingError(
                f"输出文件在处理期间出现，未执行覆盖：{output_path}\n"
                "如确认覆盖，请显式添加 --overwrite。"
            )
        os.replace(temporary_path, output_path)
        temporary_path = None
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def standardize_workbook(
    input_path: Path,
    output_path: Path,
    overwrite: bool,
) -> tuple[list[SheetStats], list[str]]:
    if not input_path.is_file():
        raise ProcessingError(f"输入文件不存在：{input_path}")
    if input_path.suffix.casefold() != ".xlsx":
        raise ProcessingError("输入文件必须是 .xlsx 格式。")
    if output_path.suffix.casefold() != ".xlsx":
        raise ProcessingError("输出文件必须使用 .xlsx 扩展名。")
    if _same_path(input_path, output_path):
        raise ProcessingError("输出路径不能与输入路径相同，源工作簿不会被覆盖。")
    if output_path.exists() and not overwrite:
        raise ProcessingError(
            f"输出文件已存在：{output_path}\n"
            "如确认覆盖，请显式添加 --overwrite。"
        )

    with timed_stage(f"读取 Excel：{input_path}"):
        source_workbook = load_workbook(input_path, data_only=True, read_only=False)
    try:
        with timed_stage("检查 Sheet 表头"):
            specs, skipped = preflight_workbook(source_workbook)
        output_workbook = Workbook()
        try:
            output_workbook.remove(output_workbook.active)
            smiles_cache: dict[str, SmilesAnalysis] = {}
            stats: list[SheetStats] = []

            for sheet_index, spec in enumerate(specs, start=1):
                print(
                    f"\n[Sheet {sheet_index}/{len(specs)}] {spec.name}",
                    file=sys.stderr, flush=True,
                )
                source_sheet = source_workbook[spec.name]
                output_sheet = output_workbook.create_sheet(spec.name)
                stats.append(
                    process_sheet(source_sheet, output_sheet, spec, smiles_cache)
                )

            with timed_stage(f"保存 Excel：{output_path}"):
                save_workbook_atomically(output_workbook, output_path, overwrite)
            return stats, skipped
        finally:
            output_workbook.close()
    finally:
        source_workbook.close()


def print_summary(
    stats: Sequence[SheetStats],
    skipped_sheets: Sequence[str],
    output_path: Path,
) -> None:
    for sheet_name in skipped_sheets:
        print(
            f"[跳过] Sheet {sheet_name!r}: 前 {HEADER_SCAN_ROWS} 行未找到 'MIC Value'。"
        )
    for item in stats:
        print(
            f"[{item.name}] 原始行={item.source_rows}, "
            f"删除三项全空={item.removed_blank}, 输出={item.output_rows}, "
            f"成功计算={item.calculated}, 关键词归类={item.keyword_rows}, "
            f"仅占位码删除={item.excluded_keyword_rows}, "
            f"评分使用范围中值={item.range_rows}, 待人工判断（黄色）={item.manual_rows}"
        )
        priorities = (1, 2, 3, 5, 6)
        print("  " + ", ".join(f"P{priority}={item.priority_counts.get(priority, 0)}" for priority in priorities))
    print(f"处理完成：{output_path.resolve(strict=False)}")


def parse_arguments(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="按 IC50、抑制率、MIC 的既定优先级对多 Sheet 数据分配 Activity Level。"
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=DEFAULT_INPUT,
        help=f"输入工作簿路径（默认：{DEFAULT_INPUT}）",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"输出工作簿路径（默认：{DEFAULT_OUTPUT}）",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="显式允许覆盖已存在的输出工作簿；永远不会覆盖输入文件。",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    arguments = parse_arguments(argv)
    RDLogger.DisableLog("rdApp.error")
    try:
        stats, skipped = standardize_workbook(
            arguments.input,
            arguments.output,
            arguments.overwrite,
        )
    except (ProcessingError, OSError, ValueError) as error:
        print(f"错误：{error}", file=sys.stderr)
        return 1

    print_summary(stats, skipped, arguments.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

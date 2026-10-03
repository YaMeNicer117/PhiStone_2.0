import os
import glob
import copy
from openpyxl import load_workbook

# ================= 1. 目录配置 =================
EXCEL_DIR = "EXCEL_DIR"              # 原始表格目录
OUTPUT_DIR = "OUTPUT_EXCEL_DIR"      # 修复后生成的碎片表格目录
FINAL_MERGED_NAME = "Final_Merged_Repaired.xlsx" # 最终合成的文件名

def get_column_indices(sheet):
    """获取表头对应列的 1-based 索引"""
    col_idx = {}
    for cell in sheet[1]:
        if cell.value:
            col_idx[str(cell.value).strip()] = cell.column
    return col_idx

def main():
    print(">>> 开始执行: 修复结果增量完美拼装程序 <<<")
    
    final_path = os.path.join(OUTPUT_DIR, FINAL_MERGED_NAME)

    # 1. 确定底板文件 (支持增量合并)
    if os.path.exists(final_path):
        print(f"[*] 发现已存在的合并文件，将以此为底板进行【增量合并】: {final_path}")
        base_file = final_path
    else:
        orig_files = glob.glob(os.path.join(EXCEL_DIR, "*.xlsx"))
        if not orig_files:
            print(f"[-] 错误: 在 {EXCEL_DIR} 未找到原始 Excel 文件，且无已存在的合并文件。")
            return
        base_file = orig_files[0]
        print(f"[*] 未发现历史合并记录，采用原始 Excel 为底板: {base_file}")

    # 2. 寻找所有修复后的碎片 Excel (排除最终合并文件本身，防止自我读取)
    all_out_files = glob.glob(os.path.join(OUTPUT_DIR, "*.xlsx"))
    repaired_files = [f for f in all_out_files if os.path.basename(f) != FINAL_MERGED_NAME]
    
    if not repaired_files:
        print(f"[-] 警告: 在 {OUTPUT_DIR} 未找到任何待合并的碎片 Excel 文件。")
        return
    print(f"[*] 发现 {len(repaired_files)} 个修复碎片文件，开始提取补丁数据...")

    # 3. 提取所有补丁数据，按 Source Filename 分组存入队列
    # 数据结构: { "文献名.pdf": [ {row_data1}, {row_data2} ] }
    patch_database = {}
    total_patches = 0

    for rep_file in repaired_files:
        wb_rep = load_workbook(rep_file, data_only=True)
        ws_rep = wb_rep.active
        cols = get_column_indices(ws_rep)
        
        # 必须确保碎片文件里有这几列
        if "Source Filename" not in cols or "SMILES" not in cols:
            continue
            
        col_src = cols["Source Filename"]
        col_smi = cols["SMILES"]
        col_smi_source = cols.get("SMILES Source")

        for row in range(2, ws_rep.max_row + 1):
            pdf_cell = ws_rep.cell(row=row, column=col_src)
            pdf_name = str(pdf_cell.value).strip() if pdf_cell.value else ""
            
            if not pdf_name or pdf_name.lower() in ["nan", "none", ""]:
                continue
                
            smi_cell = ws_rep.cell(row=row, column=col_smi)
            src_val = ws_rep.cell(row=row, column=col_smi_source).value if col_smi_source else ""

            patch_data = {
                "smiles": smi_cell.value,
                "smiles_source": src_val,
                "smiles_fill": copy.copy(smi_cell.fill) if smi_cell.fill and smi_cell.fill.fill_type else None,
                "pdf_fill": copy.copy(pdf_cell.fill) if pdf_cell.fill and pdf_cell.fill.fill_type else None
            }

            if pdf_name not in patch_database:
                patch_database[pdf_name] = []
            patch_database[pdf_name].append(patch_data)
            total_patches += 1

    print(f"[+] 补丁数据提取完毕。当前碎片库中共包含 {total_patches} 条记录。")

    # 4. 打开底板 Excel，逐行进行“精准注射”
    print(f"[*] 正在将补丁注射回底板表格 (完全保持原顺序与 Sheet 结构)...")
    wb_orig = load_workbook(base_file)
    
    injected_count = 0

    for sheet in wb_orig.worksheets:
        cols = get_column_indices(sheet)
        if "Source Filename" not in cols or "SMILES" not in cols:
            continue
            
        col_src = cols["Source Filename"]
        col_smi = cols["SMILES"]
        
        # 如果原表没有 SMILES Source 这列，我们就在最后一列加上它
        if "SMILES Source" not in cols:
            col_smi_source = sheet.max_column + 1
            sheet.cell(row=1, column=col_smi_source).value = "SMILES Source"
        else:
            col_smi_source = cols["SMILES Source"]

        for row in range(2, sheet.max_row + 1):
            pdf_cell = sheet.cell(row=row, column=col_src)
            pdf_name = str(pdf_cell.value).strip() if pdf_cell.value else ""
            
            # 如果是空行或者没有文献名的行，跳过
            if not pdf_name or pdf_name.lower() in ["nan", "none", ""]:
                continue

            # 如果我们在补丁库中找到了这篇文献，并且队列里还有数据
            if pdf_name in patch_database and len(patch_database[pdf_name]) > 0:
                # 弹出队列顶部的第一个补丁 (这保证了完美的顺序对应)
                patch = patch_database[pdf_name].pop(0)
                
                # 注射 SMILES 及其背景色 (黄/橙)
                smi_cell = sheet.cell(row=row, column=col_smi)
                smi_cell.value = patch["smiles"]
                if patch["smiles_fill"]:
                    smi_cell.fill = patch["smiles_fill"]
                    
                # 注射 文献名的背景色 (蓝)
                if patch["pdf_fill"]:
                    pdf_cell.fill = patch["pdf_fill"]
                    
                # 注射 SMILES Source
                src_cell = sheet.cell(row=row, column=col_smi_source)
                src_cell.value = patch["smiles_source"]

                injected_count += 1

    # 5. 保存最终文件
    wb_orig.save(final_path)
    
    print(f"\n[+] 增量拼装完成！")
    print(f"[!] 本次共刷新/覆盖了 {injected_count} 行记录的修正数据。")
    print(f"[!] 最终文件已保存至: {final_path}")

if __name__ == "__main__":
    main()